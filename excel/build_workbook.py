"""Build the Excel Windows template from the two workbook exports."""

import csv
from pathlib import Path

from openpyxl import load_workbook
import xlsxwriter


ROOT = Path(__file__).resolve().parents[1]
DESTINATION = Path(__file__).with_name("Maintenance_Planning.xlsx")
PLAN_HEADERS = [
    "Selected line", "Maintenance item", "Maintenance Plan", "Maintenance strategy",
    "Maintenance item description", "MntPlan Call No.", "Scheduled start date",
    "Order", "Completion date",
]
EXECUTION_HEADERS = [
    "Order", "Description", "Functional Location", "Main work center",
    "Basic start date", "Priority", "Maintenance Plan", "User Status",
    "Entered by", "Completion Date", "Long text1", "Long text2",
]
ANALYSIS_HEADERS = [
    "Machine", "Activity", "Maintenance Plan", "Frequency", "Last completion",
    "Next due", "Days overdue", "Annual target", "Elapsed periods", "Completed periods",
    "Remaining", "Annual compliance", "Status", "Execution records", "Plan calls",
]
DMS_HEADERS = ["Title", "Machine", "Date added", "Fix Description", "Fixed Date", "Image", "Status", "Engineering Notes", "Health and Safety", "Engineer"]


def input_rows(path, headers):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    columns = {name: index for index, name in enumerate(next(sheet.values))}
    missing = set(headers) - set(columns)
    if missing:
        raise ValueError(f"{path.name}: missing columns: {', '.join(sorted(missing))}")
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = [values[columns[name]] for name in headers]
        if any(value is not None for value in row):
            rows.append(row)
    workbook.close()
    return rows


def main():
    plan = [row for row in input_rows(ROOT / "MAINTENANCE PLANS WITH ORDERS.XLSX", PLAN_HEADERS) if not str(row[4] or "").strip().upper().startswith("COR -")]
    execution_path = ROOT / "MAINTENANCE PLANS EXECUTION.XLSX"
    execution = input_rows(execution_path, EXECUTION_HEADERS) if execution_path.exists() else []
    dms = []
    dms_path = ROOT / "DMS List (5).csv"
    if dms_path.exists():
        with dms_path.open(encoding="utf-8-sig", newline="") as source:
            reader = csv.DictReader(source)
            dms = [[record.get(header, "") for header in DMS_HEADERS] for record in reader if record.get("Title", "").strip()]
    workbook = xlsxwriter.Workbook(DESTINATION)
    workbook.set_properties({"title": "Maintenance Planning | Excel", "comments": "Paste exports in Plan and Execution, then run UpdateMaintenanceDashboard."})
    dark, green, orange = "#27383C", "#237C61", "#DB803A"
    title = workbook.add_format({"font_name": "Aptos Display", "font_size": 20, "bold": True, "font_color": dark})
    sub = workbook.add_format({"font_name": "Aptos", "font_size": 10, "font_color": "#68777D"})
    label = workbook.add_format({"font_name": "Aptos", "bold": True, "font_color": dark, "bg_color": "#EEF2F1", "bottom": 1, "bottom_color": "#CBD7D3"})
    section = workbook.add_format({"font_name": "Aptos", "bold": True, "font_color": "#FFFFFF", "bg_color": green, "font_size": 11})
    value = workbook.add_format({"font_name": "Aptos Display", "font_size": 21, "bold": True, "font_color": dark, "num_format": "0"})
    percentage = workbook.add_format({"font_name": "Aptos Display", "font_size": 21, "bold": True, "font_color": green, "num_format": "0%"})
    input_style = workbook.add_format({"font_name": "Aptos", "font_color": dark, "bg_color": "#FFF5E7", "bold": True})
    input_date_style = workbook.add_format({"font_name": "Aptos", "font_color": dark, "bg_color": "#FFF5E7", "bold": True, "num_format": "dd/mm/yyyy"})
    date_style = workbook.add_format({"num_format": "dd/mm/yyyy"})
    note = workbook.add_format({"font_name": "Aptos", "font_color": "#53646A", "text_wrap": True, "valign": "vcenter"})

    config = workbook.add_worksheet("Settings")
    config.set_tab_color(orange)
    config.set_column("A:A", 25)
    config.set_column("B:B", 22)
    config.set_column("C:C", 75)
    config.write("A1", "MAINTENANCE PLANNING", title)
    config.write("A3", "Report year", label)
    config.write_formula("B3", "=YEAR(TODAY())", input_style, 2026)
    config.write("A4", "Reference date", label)
    config.write_formula("B4", "=TODAY()", input_date_style)
    config.write("C4", "Change the date for historical snapshots; paste a date as a value to override TODAY().", note)
    config.write("A5", "Machine filter", label)
    config.write("B5", "All", input_style)
    config.write("C5", "Enter All or a machine name exactly as shown in Analysis.", note)
    config.write("A6", "Week starting", label)
    config.write_formula("B6", "=TODAY()-WEEKDAY(TODAY(),2)+1", input_date_style)
    config.write("C6", "Monday of the week shown in Weekly Report.", note)
    config.write("A7", "DMS machine", label)
    config.write("B7", "All", input_style)
    config.write("C7", "Enter All or a DMS machine name for the action board.", note)
    config.data_validation("B3", {"validate": "integer", "criteria": "between", "minimum": 2000, "maximum": 2100})
    config.write("A8", "Refresh", section)
    config.merge_range("A9:C10", "For each new export, clear the old input sheet with Ctrl+A, Delete before pasting the complete file into A1, including headers. Run UpdateMaintenanceDashboard (Alt+F8) to rebuild all reports. Save as .xlsm after importing the VBA module.", note)

    guide = workbook.add_worksheet("Start Here")
    guide.set_tab_color(orange)
    guide.set_column("A:A", 5)
    guide.set_column("B:B", 105)
    guide.set_default_row(22)
    guide.write("B2", "MAINTENANCE PLANNING / EXCEL", title)
    for row, text in enumerate([
        "1. Open Maintenance_Planning.xlsx in Microsoft Excel for Windows (desktop).",
        "2. Press Alt+F11, choose File > Import File, and select MaintenancePlanning.bas from the same folder.",
        "3. Save as Excel Macro-Enabled Workbook (*.xlsm). Enable content only if you trust this local module.",
        "4. In each input sheet (Plan, Execution, DMS Input): Ctrl+A, Delete, then paste the entire new export into A1 including headers.",
        "5. Adjust year, reference date and week in Settings. Run UpdateMaintenanceDashboard (Alt+F8).",
        "6. Use the filter in Analysis or type a machine in Settings!B5 and refresh for the priority list.",
        "The annual target uses the plan frequency. Completed periods use Completion Date in the report year,",
        "deduplicated by Basic start period and capped at the periods elapsed by the reference date.",
        "The latest completion drives the next due date and overdue status, independently of annual compliance.",
        "DMS Board uses the pasted action list; Weekly Report lists activities with a due date in the chosen week.",
        "This .xlsx template has no embedded macro: import MaintenancePlanning.bas once and save as .xlsm.",
    ], 4):
        guide.write(row, 1, text, note)
    guide.set_row(13, 35)

    for name, headers, rows, date_cols in [
        ("Plan", PLAN_HEADERS, plan, {6, 8}),
        ("Execution", EXECUTION_HEADERS, execution, {4, 9}),
    ]:
        sheet = workbook.add_worksheet(name)
        sheet.set_tab_color(orange if name == "Plan" else green)
        sheet.freeze_panes(1, 0)
        sheet.set_column(0, len(headers) - 1, 19)
        sheet.set_column(1 if name == "Execution" else 4, 1 if name == "Execution" else 4, 40)
        if name == "Execution":
            sheet.set_column(10, 11, 45)
        for col, header in enumerate(headers):
            sheet.write(0, col, header, label)
        for index, row in enumerate(rows, 1):
            for col, cell in enumerate(row):
                if cell is None:
                    continue
                if col in date_cols and hasattr(cell, "year"):
                    sheet.write_datetime(index, col, cell, date_style)
                elif isinstance(cell, (float, int)):
                    sheet.write_number(index, col, cell)
                else:
                    sheet.write_string(index, col, str(cell))
        sheet.autofilter(0, 0, max(1, len(rows)), len(headers) - 1)

    dms_input = workbook.add_worksheet("DMS Input")
    dms_input.set_tab_color(orange)
    dms_input.freeze_panes(1, 0)
    dms_input.set_column("A:A", 55)
    dms_input.set_column("B:C", 19)
    dms_input.set_column("D:D", 45)
    dms_input.set_column("E:G", 18)
    dms_input.set_column("H:H", 45)
    dms_input.set_column("I:J", 20)
    for index, heading in enumerate(DMS_HEADERS):
        dms_input.write(0, index, heading, label)
    for index, row in enumerate(dms, 1):
        for col, cell in enumerate(row):
            if cell:
                dms_input.write_string(index, col, cell)
    dms_input.autofilter(0, 0, max(1, len(dms)), len(DMS_HEADERS) - 1)

    analysis = workbook.add_worksheet("Analysis")
    analysis.set_tab_color(green)
    analysis.freeze_panes(1, 3)
    analysis.set_column("A:A", 22)
    analysis.set_column("B:B", 52)
    analysis.set_column("C:D", 21)
    analysis.set_column("E:F", 18, date_style)
    analysis.set_column("G:O", 18)
    for col, header in enumerate(ANALYSIS_HEADERS):
        analysis.write(0, col, header, label)
    analysis.conditional_format("G2:G5000", {"type": "cell", "criteria": ">", "value": 0, "format": workbook.add_format({"bg_color": "#FBE7E2", "font_color": "#B54940"})})
    analysis.conditional_format("L2:L5000", {"type": "3_color_scale", "min_color": "#F2B3A6", "mid_color": "#F6E9B7", "max_color": "#82C7A4"})

    board = workbook.add_worksheet("Dashboard")
    board.set_tab_color(green)
    board.hide_gridlines(2)
    board.set_zoom(85)
    board.set_column("A:A", 4)
    board.set_column("B:B", 27)
    board.set_column("C:I", 17)
    board.set_column("J:J", 4)
    board.merge_range("B2:I2", "MAINTENANCE PLAN DASHBOARD", title)
    board.merge_range("B3:I3", "Live plan control / annual compliance and next due dates", sub)
    for col, heading in [(1, "ACTIVITIES"), (3, "ANNUAL COMPLIANCE"), (5, "OVERDUE"), (7, "DUE SOON")]:
        board.write(4, col, heading, section)
    board.write("B6", 0, value)
    board.write("D6", 0, percentage)
    board.write("F6", 0, value)
    board.write("H6", 0, value)
    board.merge_range("B9:I9", "PRIORITIZED ACTIVITIES", section)
    priority_headers = ["Machine", "Activity", "Frequency", "Last completed", "Next due", "Late (days)", "Annual periods", "Status"]
    for col, heading in enumerate(priority_headers, 1):
        board.write(9, col, heading, label)
    board.merge_range("B34:I34", "COMPLIANCE BY MACHINE", section)
    for col, heading in enumerate(["Machine", "Completed", "Expected", "Compliance"], 1):
        board.write(34, col, heading, label)
    chart = workbook.add_chart({"type": "bar"})
    chart.add_series({"name": "Annual compliance", "categories": "=Dashboard!$B$36:$B$55", "values": "=Dashboard!$E$36:$E$55", "fill": {"color": green}, "border": {"none": True}})
    chart.set_title({"name": "Annual compliance by machine"})
    chart.set_legend({"none": True})
    chart.set_x_axis({"num_format": "0%", "min": 0, "max": 1})
    chart.set_chartarea({"border": {"none": True}})
    board.insert_chart("G35", chart, {"x_scale": 1.55, "y_scale": 1.25})

    weekly = workbook.add_worksheet("Weekly Report")
    weekly.set_tab_color(green)
    weekly.hide_gridlines(2)
    weekly.freeze_panes(6, 2)
    weekly.set_column("A:A", 4)
    weekly.set_column("B:B", 20)
    weekly.set_column("C:C", 55)
    weekly.set_column("D:E", 18, date_style)
    weekly.set_column("F:F", 18)
    weekly.set_column("G:G", 22)
    weekly.merge_range("B2:G2", "WEEKLY MAINTENANCE REPORT", title)
    weekly.write_formula("B4", "=Settings!B6", date_style)
    weekly.write("C4", "Week start / Monday", sub)
    for col, heading in enumerate(["Machine", "Activity", "Next due", "Last completed", "Frequency", "Status"], 1):
        weekly.write(5, col, heading, label)

    dms_board = workbook.add_worksheet("DMS Board")
    dms_board.set_tab_color(green)
    dms_board.hide_gridlines(2)
    dms_board.freeze_panes(10, 3)
    dms_board.set_column("A:A", 4)
    dms_board.set_column("B:B", 23)
    dms_board.set_column("C:C", 54)
    dms_board.set_column("D:D", 20, date_style)
    dms_board.set_column("E:E", 17)
    dms_board.set_column("F:F", 23)
    dms_board.set_column("G:G", 55)
    dms_board.set_column("H:I", 22)
    dms_board.merge_range("B2:I2", "DMS BOARD REPORT", title)
    dms_board.merge_range("B3:I3", "Action status / line priorities / days open", sub)
    for col, heading in [(1, "TOTAL ACTIONS"), (3, "COMPLETED"), (5, "OPEN / PENDING"), (7, "COMPLETION RATE")]:
        dms_board.write(4, col, heading, section)
    for cell, number_format in [("B6", value), ("D6", value), ("F6", value), ("H6", percentage)]:
        dms_board.write(cell, 0, number_format)
    for col, heading in enumerate(["Machine", "Activity", "Date added", "Days open", "Status", "Fix / Engineering notes", "H&S", "Engineer"], 1):
        dms_board.write(9, col, heading, label)
    dms_board.conditional_format("E11:E5010", {"type": "cell", "criteria": ">", "value": 14, "format": workbook.add_format({"font_color": "#B54940", "bg_color": "#FBE7E2"})})
    workbook.close()
    print(f"Created {DESTINATION} with {len(plan)} plan rows, {len(execution)} execution rows, and {len(dms)} DMS actions")


if __name__ == "__main__":
    main()